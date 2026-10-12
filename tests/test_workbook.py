"""Verification workbook: the Excel formulas must recompute exactly what InnoPetrol computed.

The formulas are evaluated with pycel, a spreadsheet engine independent of InnoPetrol.
"""

import numpy as np
import pytest
from openpyxl import load_workbook
from pycel import ExcelCompiler

from innopetrol.model import petro
from innopetrol.model.survey import Survey
from innopetrol.model.tops import Top
from innopetrol.model.well import Curve
from innopetrol.model.workbook import verification_workbook, write_table

BASE = {"gr_curve": "SGR", "gr_clean": 20, "gr_shale": 120, "rhob_curve": "RHOB", "rt_curve": "LLD"}

CASES = {
    "archie, density": {},
    "larionov, neutron-density, simandoux, zones, gradient": {
        "vsh_method": "larionov_tertiary",
        "nphi_curve": "TNPH",
        "phi_method": "nd_average",
        "sw_model": "simandoux",
        "rsh": 3.0,
        "temp_mode": "gradient",
        "temp_unit": "degF",
        "surface_temp": 60,
        "gradient": 1.5,
        "rw_temp": 75,
        "zones": {"B": {"rw": 0.03, "m": 1.8, "gr_shale": 110}},
    },
    "older rocks, RMS, indonesia, temperature curve": {
        "vsh_method": "larionov_older",
        "nphi_curve": "TNPH",
        "phi_method": "nd_rms",
        "sw_model": "indonesia",
        "temp_mode": "curve",
        "temp_curve": "BHT",
        "temp_unit": "degF",
        "rw_temp": 75,
    },
}


def column(ws, header_start: str) -> int:
    return next(c.column for c in ws[1] if str(c.value).startswith(header_start))


@pytest.mark.parametrize("case", list(CASES))
def test_formulas_recompute_innopetrol(testdata_well, tmp_path, case):
    w = testdata_well
    w.tops = [Top("A", 7700), Top("B", 8000)]
    w.survey = Survey(np.array([0.0, 7000, 9000]), np.array([0.0, 15, 30]), np.array([0.0, 45, 60]))
    w.add_curve(Curve("BHT", "DEGC", "", np.linspace(60, 90, w.depth.size)))
    p = petro.PetroParams(**{**BASE, **CASES[case]})
    path = tmp_path / "check.xlsx"
    n = verification_workbook(path, w, p, 7600, 8200)
    assert n > 100
    ws = load_workbook(path)["Data"]
    excel = ExcelCompiler(filename=str(path))
    rows = np.linspace(2, n + 1, 15).astype(int)
    checked = 0
    for name in ("VSH", "PHIE", "SW"):
        col_x = column(ws, {"VSH": "Vsh (Excel)", "PHIE": "PHIE (Excel)", "SW": "Sw (Excel)"}[name])
        col_p = column(ws, f"{name} InnoPetrol")
        for r in rows:
            mine = ws.cell(row=r, column=col_p).value
            theirs = excel.evaluate(f"Data!{ws.cell(row=r, column=col_x).coordinate}")
            if mine is None:
                assert theirs in ("", None) or not np.isfinite(theirs)
                continue
            assert theirs == pytest.approx(mine, abs=1e-9), (name, r)
            checked += 1
    assert checked > 30


def test_zone_parameters_are_written_per_sample(testdata_well, tmp_path):
    w = testdata_well
    w.tops = [Top("B", 8000)]
    p = petro.PetroParams(**BASE, zones={"B": {"rw": 0.03}})
    path = tmp_path / "z.xlsx"
    verification_workbook(path, w, p)
    ws = load_workbook(path)["Data"]
    zone, rw = column(ws, "Zone"), next(c.column for c in ws[1] if c.value == "Rw")
    values = {
        (ws.cell(row=r, column=zone).value or "", ws.cell(row=r, column=rw).value) for r in range(2, ws.max_row + 1)
    }
    assert values == {("", 0.05), ("B", 0.03)}


def test_table_workbook(tmp_path):
    path = tmp_path / "t.xlsx"
    write_table(path, "Zones: summary", ["Zone", "Net pay"], [["A", 12.5], ["B", float("nan")]])
    ws = load_workbook(path).active
    assert ws.title == "Zones summary" and ws["B2"].value == 12.5 and ws["B3"].value is None
