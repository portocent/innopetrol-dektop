"""Excel workbooks (.xlsx): tables, and a verification workbook for the petrophysics.

The verification workbook lets an engineer check InnoPetrol with an independent implementation: their
spreadsheet. For every sample it writes the inputs (in the units the equations use) and the
parameters of the sample's zone, then the same equations as **Excel formulas**, InnoPetrol's results
and the difference. Formulas are stored in Excel's English notation, which Excel and LibreOffice show
in the user's language, so the file works everywhere (a CSV with formulas would not).
"""

from __future__ import annotations

from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from innopetrol import __version__
from innopetrol.i18n import _
from innopetrol.model import petro, units
from innopetrol.model.tops import zones_from_tops

BOLD = Font(bold=True)
INPUT_FILL = PatternFill("solid", fgColor="EAF2F8")
FORMULA_FILL = PatternFill("solid", fgColor="FEF9E7")
RESULT_FILL = PatternFill("solid", fgColor="E9F7EF")
DIFF_FILL = PatternFill("solid", fgColor="FDEDEC")


def write_table(path: str | Path, title: str, headers: list[str], rows: list[list]) -> None:
    """A single-sheet workbook with a bold header row (numbers stay numbers)."""
    wb = Workbook()
    ws = wb.active
    ws.title = _sheet_name(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = BOLD
    for row in rows:
        ws.append([_cell_value(v) for v in row])
    ws.freeze_panes = "A2"
    _autosize(ws)
    wb.save(path)


# ---------------------------------------------------------------- verification workbook


@dataclass
class _Column:
    name: str
    header: str
    kind: str  # "input" | "formula" | "result" | "diff"
    value: Callable[[int, int], object]  # (sample index, sheet row) -> value or formula


class _Sheet:
    """Columns of the data sheet; formulas refer to other columns by name: {GR} → "N5"."""

    def __init__(self):
        self.columns: list[_Column] = []

    def letter(self, name: str) -> str:
        return get_column_letter(next(i for i, c in enumerate(self.columns, 1) if c.name == name))

    def add(self, name: str, header: str, kind: str, value) -> None:
        self.columns.append(_Column(name, header, kind, value))

    def formula(self, name: str, header: str, template: str) -> None:
        """A formula column; `template` uses {NAME} for "the cell of column NAME in this row"."""

        def value(_i: int, row: int) -> str:
            refs = {c.name: f"{self.letter(c.name)}{row}" for c in self.columns}
            return "=" + template.format(**refs)

        self.add(name, header, "formula", value)

    def has(self, name: str) -> bool:
        return any(c.name == name for c in self.columns)


def _clip(expr: str, lo: float, hi: float) -> str:
    return f"MAX({lo},MIN({hi},{expr}))"


def _zone_names(well) -> np.ndarray:
    names = np.full(well.depth.shape, "", dtype=object)
    for z in zones_from_tops(well.tops, well.stop):
        names[well.interval_mask(z.top, z.base)] = z.name
    return names


def _input(well, mnemonic: str | None, unit: str) -> np.ndarray | None:
    """A curve's values in the unit the equations use (None if missing or of another kind of unit)."""
    from innopetrol.model.mnemonics import data_unit

    if mnemonic is None or mnemonic not in well.curves:
        return None
    curve = well.curves[mnemonic]
    try:
        return units.convert(curve.values, data_unit(curve, unit) or unit, unit)
    except ValueError:
        return None


@dataclass
class _Inputs:
    """Input arrays of the calculation, in the units the equations use (None = step not computed)."""

    results: dict[str, np.ndarray]
    zones: np.ndarray
    gr: np.ndarray | None
    rhob: np.ndarray | None
    nphi: np.ndarray | None
    rt: np.ndarray | None
    temperature: np.ndarray | None
    tvd: np.ndarray | None

    def required(self) -> list[np.ndarray]:
        return [a for a in (self.gr, self.rhob, self.nphi, self.rt, self.temperature) if a is not None]


def _gather(well, p: petro.PetroParams) -> _Inputs:
    results = {c.mnemonic: c.values for c in petro.compute(well, p)}
    temperature = None
    if "RWT" in results:
        temperature = results.get("TFORM") if p.temp_mode == "gradient" else _input(well, p.temp_curve, p.temp_unit)
    return _Inputs(
        results=results,
        zones=_zone_names(well),
        gr=well.curves[p.gr_curve].values if "VSH" in results and p.gr_curve else None,  # used as written
        rhob=_input(well, p.rhob_curve, "g/cm3") if "PHIT" in results else None,
        nphi=_input(well, p.nphi_curve, "v/v") if "PHID" in results else None,
        rt=well.curves[p.rt_curve].values if "SW" in results and p.rt_curve else None,
        temperature=temperature,
        tvd=well.tvd() if p.temp_mode == "gradient" and "RWT" in results else None,
    )


def verification_workbook(
    path: str | Path,
    well,
    p: petro.PetroParams,
    top: float | None = None,
    base: float | None = None,
    max_rows: int = 20000,
) -> int:
    """Writes the verification workbook for an interval. Returns the number of samples written.

    Only samples where every input is known are written, so each row checks a complete calculation.
    """
    data = _gather(well, p)
    required = data.required()
    if not required:
        return 0
    mask = well.interval_mask(top, base) & np.all([np.isfinite(a) for a in required], axis=0)
    idx = np.flatnonzero(mask)[:max_rows]
    sheet = _Sheet()
    _add_inputs(sheet, well, p, data)
    _add_formulas(
        sheet, p, has_vsh=data.gr is not None, temperature=data.temperature is not None, gradient=data.tvd is not None
    )
    _add_results(sheet, data.results)
    wb = Workbook()
    _readme(wb.active, well, p, len(idx), top, base)
    _write_data(wb.create_sheet(_("Data")), sheet, idx)
    wb.save(path)
    return len(idx)


def _add_inputs(sheet: _Sheet, well, p: petro.PetroParams, data: _Inputs) -> None:
    """Inputs and, for every sample, the parameters of its zone."""
    zones = data.zones
    sheet.add("DEPTH", _("Depth MD ({unit})").format(unit=well.depth_unit), "input", lambda i, r: well.depth[i])
    sheet.add("ZONE", _("Zone"), "input", lambda i, r: zones[i])

    def param(key: str):
        return lambda i, r: p.value(key, str(zones[i]) or None)

    def column(name: str, header: str, values: np.ndarray) -> None:
        sheet.add(name, header, "input", lambda i, r: values[i])

    if data.gr is not None:
        column("GR", f"GR ({well.curves[p.gr_curve].unit})", data.gr)
        sheet.add("GRC", _("Clean GR"), "input", param("gr_clean"))
        sheet.add("GRS", _("Shale GR"), "input", param("gr_shale"))
    if data.rhob is not None:
        column("RHOB", "RHOB (g/cm³)", data.rhob)
        sheet.add("RHOMA", _("ρ matrix"), "input", param("rho_ma"))
        sheet.add("RHOFL", _("ρ fluid"), "input", param("rho_fl"))
    if data.nphi is not None:
        column("NPHI", "NPHI (v/v)", data.nphi)
    if data.rt is not None:
        column("RT", f"Rt ({well.curves[p.rt_curve].unit})", data.rt)
        for key, label in (("rw", "Rw"), ("a", "a"), ("m", "m"), ("n", "n"), ("rsh", "Rsh")):
            sheet.add(key.upper(), label, "input", param(key))
    if data.temperature is not None:
        unit = units.label(p.temp_unit)
        if data.tvd is not None:
            column("TVD", _("TVD ({unit})").format(unit=well.depth_unit), data.tvd)
            sheet.add(
                "TSURF", _("Surface temperature ({unit})").format(unit=unit), "input", lambda i, r: p.surface_temp
            )
            sheet.add("GRAD", _("Gradient ({unit}/100)").format(unit=unit), "input", lambda i, r: p.gradient)
        else:
            column("TCURVE", _("Temperature ({unit})").format(unit=unit), data.temperature)
        sheet.add("TRW", _("Rw measured at ({unit})").format(unit=unit), "input", lambda i, r: p.rw_temp)


def _add_results(sheet: _Sheet, results: dict[str, np.ndarray]) -> None:
    """InnoPetrol's results next to the spreadsheet's, and their difference."""
    for name in ("VSH", "PHIE", "SW"):
        if name in results and sheet.has(name + "_X"):
            values = results[name]
            sheet.add(name, _("{curve} InnoPetrol").format(curve=name), "result", lambda i, r, v=values: v[i])
            template = f'IF(AND(ISNUMBER({{{name}_X}}),ISNUMBER({{{name}}})),ABS({{{name}_X}}-{{{name}}}),"")'
            sheet.formula("D" + name, _("Difference {curve}").format(curve=name), template)
            sheet.columns[-1].kind = "diff"


def _write_data(ws, sheet: _Sheet, idx: np.ndarray) -> None:
    fills = {"input": INPUT_FILL, "formula": FORMULA_FILL, "result": RESULT_FILL, "diff": DIFF_FILL}
    ws.append([c.header for c in sheet.columns])
    for col, c in enumerate(sheet.columns, 1):
        cell = ws.cell(row=1, column=col)
        cell.font, cell.fill = BOLD, fills[c.kind]
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row, i in enumerate(idx, 2):
        ws.append([_cell_value(c.value(int(i), row)) for c in sheet.columns])
    ws.freeze_panes = "C2"
    _autosize(ws, maximum=16)


def _add_formulas(sheet: _Sheet, p: petro.PetroParams, has_vsh: bool, temperature: bool, gradient: bool) -> None:
    """The petrophysics equations as Excel formulas (same order and limits as InnoPetrol)."""
    if has_vsh:
        sheet.formula("IGR_X", "IGR (Excel)", _clip("({GR}-{GRC})/({GRS}-{GRC})", 0, 1))
        vsh = {
            "linear": "{IGR_X}",
            "larionov_tertiary": _clip("0.083*(2^(3.7*{IGR_X})-1)", 0, 1),
            "larionov_older": _clip("0.33*(2^(2*{IGR_X})-1)", 0, 1),
        }[p.vsh_method]
        sheet.formula("VSH_X", "Vsh (Excel)", vsh)
    if sheet.has("RHOB"):
        phid = "({RHOMA}-{RHOB})/({RHOMA}-{RHOFL})"
        if sheet.has("NPHI"):
            sheet.formula("PHID_X", "PHID (Excel)", _clip(phid, -0.15, 0.6))
            phit = "({PHID_X}+{NPHI})/2" if p.phi_method == "nd_average" else "SQRT(({PHID_X}^2+{NPHI}^2)/2)"
        else:
            phit = phid
        sheet.formula("PHIT_X", "PHIT (Excel)", _clip(phit, 0, 0.6))
        sheet.formula("PHIE_X", "PHIE (Excel)", _clip("{PHIT_X}*(1-{VSH_X})", 0, 0.6) if has_vsh else "{PHIT_X}")
    if not sheet.has("RT"):
        return
    rw = "{RW}"
    if temperature:
        if gradient:
            sheet.formula("T_X", _("Temperature (Excel)"), "{TSURF}+{GRAD}*{TVD}/100")
            t = "{T_X}"
        else:
            t = "{TCURVE}"
        k = petro.ARPS_K[p.temp_unit]
        sheet.formula("RWT_X", "Rw(T) (Excel)", f"{{RW}}*({{TRW}}+{k})/({t}+{k})")
        rw = "{RWT_X}"
    model = p.sw_model if has_vsh and p.sw_model in petro.SW_MODELS else "archie"
    if model == "simandoux":
        core = (
            f"({{A}}*{rw}/(2*{{PHIE_X}}^{{M}}))"
            f"*(SQRT(({{VSH_X}}/{{RSH}})^2+4*{{PHIE_X}}^{{M}}/({{A}}*{rw}*{{RT}}))-{{VSH_X}}/{{RSH}})"
        )
    elif model == "indonesia":
        core = (
            f"((1/SQRT({{RT}}))/({{VSH_X}}^(1-{{VSH_X}}/2)/SQRT({{RSH}})+SQRT({{PHIE_X}}^{{M}}/({{A}}*{rw}))))"
            f"^(2/{{N}})"
        )
    else:
        core = f"(({{A}}*{rw})/({{PHIE_X}}^{{M}}*{{RT}}))^(1/{{N}})"
    sheet.formula("SW_X", "Sw (Excel)", f'IF(AND({{PHIE_X}}>0,{{RT}}>0),{_clip(core, 0, 1)},"")')


def _readme(ws, well, p: petro.PetroParams, n_rows: int, top, base) -> None:
    ws.title = _("Read me")
    lines = [
        (_("Petrophysics verification workbook"), True),
        (_("Well: {name}").format(name=well.name), False),
        (
            _("Interval: {top} – {base} {unit} (MD)").format(
                top=f"{top if top is not None else well.start:g}",
                base=f"{base if base is not None else well.stop:g}",
                unit=well.depth_unit,
            ),
            False,
        ),
        (_("Samples: {count} (only those with every input known)").format(count=n_rows), False),
        (
            _("Methods: Vsh {vsh} · porosity {phi} · Sw {sw}").format(
                vsh=_(petro.VSH_METHODS[p.vsh_method]),
                phi=_(petro.PHI_METHODS[p.phi_method]),
                sw=_(petro.SW_MODELS[p.sw_model]),
            ),
            False,
        ),
        ("", False),
        (
            _(
                "Sheet “Data”: blue columns are inputs (in the units the equations use) and the parameters of "
                "each sample's zone; yellow columns recompute the results with spreadsheet formulas; green "
                "columns are InnoPetrol's results; red columns are the absolute differences."
            ),
            False,
        ),
        (_("Change any input or parameter to see the effect: the formulas recalculate."), False),
        ("", False),
        (_("Created by InnoPetrol {version}").format(version=__version__), False),
    ]
    for text, bold in lines:
        ws.append([text])
        if bold:
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 120
    for row in ws.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")


def _cell_value(v):
    if isinstance(v, (float, np.floating)):
        return float(v) if np.isfinite(v) else None
    if isinstance(v, np.integer):
        return int(v)
    return v


def _sheet_name(title: str) -> str:
    return "".join(c for c in title if c not in "[]:*?/\\")[:31] or "Sheet"


def _autosize(ws, maximum: int = 40) -> None:
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col[:50] if c.value is not None), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 8), maximum)
